"""Seed dataset tests: deterministic, sized, secret-free, offline.

The seed script is loaded by path (it is tooling, not an installed
module). Nothing here touches AWS: builders are pure functions.
"""
import importlib.util
import re
from pathlib import Path

import pytest

SEED_PATH = (Path(__file__).parent.parent / "scripts"
             / "seed_demo_workspace.py")


def _load_seed():
    spec = importlib.util.spec_from_file_location(
        "seed_demo_workspace", SEED_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def seed():
    return _load_seed()


@pytest.fixture(scope="module")
def built(seed):
    return seed.build_all()


def test_dataset_has_intended_count(built):
    assert 15 <= len(built) <= 25, f"count={len(built)}"


def test_dataset_total_size_is_meaningful_but_bounded(built):
    total = sum(len(payload) for payload in built.values())
    assert 500_000 <= total <= 2_500_000, f"total={total}"
    assert min(len(p) for p in built.values()) < 20_000
    assert max(len(p) for p in built.values()) > 70_000


def test_every_doc_fits_analysis_limit(built):
    for key, payload in built.items():
        assert len(payload) <= 100_000, f"{key} exceeds reader cap"


def _auditable_text(built):
    """Decodable text plus xlsx cell values (binary zips excluded)."""
    from openpyxl import load_workbook
    import io
    parts = []
    for key, payload in built.items():
        if key.endswith(".xlsx"):
            workbook = load_workbook(io.BytesIO(payload), read_only=True,
                                     data_only=True)
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    parts.extend(str(value) for value in row
                                 if value is not None)
        else:
            parts.append(payload.decode("utf-8", "replace"))
    return "\n".join(parts)


def test_no_secrets_or_real_pii_patterns(built):
    blob = _auditable_text(built)
    for pattern in (r"AKIA[0-9A-Z]{16}", r"gsk_[A-Za-z0-9]+", r"nvapi-",
                    r"xai-", r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
                    r"(?i)\bpassword\s*[:=]", r"(?i)\bapi[_-]?key\s*[:=]"):
        assert not re.search(pattern, blob), f"secret pattern: {pattern}"


def test_duplicate_pairs_present(seed, built):
    pairs = seed.duplicate_pairs()
    assert len(pairs) >= 1
    for first, second in pairs:
        assert first in built and second in built


def test_age_variation_present(seed, built):
    eras = {spec["era"] for spec in seed.manifest(built)}
    assert {"current", "older", "stale"} <= eras


def test_generation_is_deterministic(seed, built):
    second = seed.build_all()
    assert built.keys() == second.keys()
    for key in built:
        assert built[key] == second[key], f"nondeterministic: {key}"


def test_manifest_matches_built_keys(seed, built):
    assert sorted(spec["key"] for spec in seed.manifest(built)) == sorted(
        built)


def test_reset_only_touches_manifest_keys(seed, built):
    from unittest.mock import MagicMock
    s3 = MagicMock()
    seed.reset(s3_client=s3, bucket="bucket")
    deleted = sorted(call.kwargs["Key"]
                     for call in s3.delete_object.call_args_list)
    assert deleted == sorted(spec["key"] for spec in seed.manifest(built))
